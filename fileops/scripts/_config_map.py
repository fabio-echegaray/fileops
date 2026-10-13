import os
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from fileops.export.config import _read_cfg_file, build_config_list, find_default_files
from fileops.logger import get_logger
from fileops.scripts._config_update import rename_movie_filename
from fileops.scripts._utils import read_summary_list

log = get_logger(name='config_map')

REGISTRY_COLUMNS = ["cfg_folder", "cfg_path", "image_path", "image_series_id"]


def _empty_registry() -> pd.DataFrame:
    return pd.DataFrame(columns=REGISTRY_COLUMNS)


def _coerce_series(value) -> int:
    try:
        if value is None or value == "" or (isinstance(value, float) and np.isnan(value)):
            return 0
        return int(float(value))
    except (ValueError, TypeError):
        return 0


def _cell(value):
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    return value


def _image_key(img_path, series) -> tuple:
    return ("" if img_path is None else str(img_path), _coerce_series(series))


def relative_image_path(image_path, relative_to):
    """Render an image path the way the registry stores it.

    Paths already relative are kept as written (they are relative to whatever
    base the user works with); absolute paths are made relative to
    ``relative_to`` when one is given.
    """
    p = Path(image_path)
    if relative_to is not None and p.is_absolute():
        try:
            return Path(os.path.relpath(p, relative_to)).as_posix()
        except ValueError:
            pass
    return p.as_posix()


def build_config_map(cfg_root, relative_to=None, progress_callback=None):
    """Build the Configs registry DataFrame from the cfg files under cfg_root.

    One row per cfg file (multi-MOVIE files collapse to a single row); the
    columns are cfg_folder, cfg_path, image_path, image_series_id.
    """
    df = build_config_list(cfg_root, relative_to=relative_to,
                           progress_callback=progress_callback)
    if df is None or df.empty:
        return _empty_registry()
    df = df.rename(columns={"image_series": "image_series_id"})
    df["image_series_id"] = df["image_series_id"].apply(_coerce_series)
    df["image_path"] = df["image_path"].apply(lambda p: relative_image_path(p, relative_to))
    df = df.drop_duplicates(subset=["cfg_path"], keep="first")
    return df[REGISTRY_COLUMNS].reset_index(drop=True)


def read_config_map(content_path):
    """Read the Configs sheet of a cfg_content workbook."""
    content_path = Path(content_path)
    if not content_path.exists():
        raise FileNotFoundError(content_path)
    book = load_workbook(content_path, read_only=True)
    try:
        if "Configs" not in book.sheetnames:
            return _empty_registry()
    finally:
        book.close()
    df = pd.read_excel(content_path, sheet_name="Configs").fillna("")
    for col in REGISTRY_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    return df


def write_config_map(content_path, df_map):
    """Replace only the Configs sheet, preserving the Content sheet."""
    content_path = Path(content_path)
    content_path.parent.mkdir(parents=True, exist_ok=True)
    df = df_map.reset_index(drop=True)
    if content_path.exists():
        book = load_workbook(content_path)
        if "Configs" in book.sheetnames:
            del book["Configs"]
    else:
        book = Workbook()
        book.remove(book.active)
        book.create_sheet("Content")
    ws = book.create_sheet("Configs")
    ws.append(list(df.columns))
    for _, row in df.iterrows():
        ws.append([_cell(v) for v in row.tolist()])
    book.save(content_path)


def _row_dict(series: pd.Series) -> dict:
    out = {}
    for col in REGISTRY_COLUMNS:
        if col not in series.index:
            out[col] = ""
            continue
        value = series[col]
        out[col] = "" if pd.isna(value) else value
    return out


def sync_config_map(content_path, cfg_root, summary_path=None, relative_to=None,
                    progress_callback=None):
    """Reconcile the Configs registry with the cfg files (and optionally the
    summary media rows), preserving user cfg_folder edits."""
    cfg_root = Path(cfg_root)
    df_map = build_config_map(cfg_root, relative_to=relative_to,
                              progress_callback=progress_callback)
    if Path(content_path).exists():
        existing = read_config_map(content_path)
    else:
        existing = _empty_registry()

    rows = []
    key_to_row: dict = {}
    for _, r in existing.iterrows():
        row = _row_dict(r)
        rows.append(row)
        key_to_row[_image_key(row["image_path"], row["image_series_id"])] = row

    for _, r in df_map.iterrows():
        key = _image_key(r["image_path"], r["image_series_id"])
        row = {
            "cfg_path":        r["cfg_path"],
            "image_path":      r["image_path"],
            "image_series_id": _coerce_series(r["image_series_id"]),
        }
        prev = key_to_row.get(key)
        if prev is not None:
            prev_cfg_folder = prev.get("cfg_folder", "")
            for col, value in row.items():
                prev[col] = value
            if prev_cfg_folder and prev_cfg_folder != r["cfg_folder"]:
                prev["cfg_folder"] = prev_cfg_folder
            else:
                prev["cfg_folder"] = r["cfg_folder"]
        else:
            row["cfg_folder"] = r["cfg_folder"]
            rows.append(row)
            key_to_row[key] = row

    if summary_path is not None:
        summary = Path(summary_path)
        if summary.exists():
            dfs, _ = read_summary_list(summary)
            if len(dfs):
                for _, r in dfs.iterrows():
                    folder = r.get("folder")
                    filename = r.get("filename")
                    if not filename:
                        continue
                    folder_path = Path(folder) if folder else Path("")
                    if not folder_path.is_absolute() and relative_to is not None:
                        folder_path = Path(relative_to) / folder_path
                    img = relative_image_path((folder_path / filename).as_posix(), relative_to)
                    series = _coerce_series(r["image_series_id"]) \
                        if "image_series_id" in r.index else 0
                    key = (img, series)
                    if key in key_to_row:
                        continue
                    rows.append({
                        "cfg_folder":      "",
                        "cfg_path":        "",
                        "image_path":      img,
                        "image_series_id": series,
                    })
                    key_to_row[key] = rows[-1]

    out = pd.DataFrame(rows)
    if not out.empty:
        out = out.sort_values(["image_path", "image_series_id"]).reset_index(drop=True)
    write_config_map(content_path, out)
    return out


def rename_config_folders(content_path, cfg_root, relative_to=None,
                          progress_callback=None):
    """Rename cfg folders to match the Configs registry (git mv with -C,
    then rename the MOVIE filename field), updating cfg_path in the registry."""
    df = read_config_map(content_path)
    dfo = df.copy()
    total = int((dfo["cfg_path"].astype(str).str.strip() != "").sum())
    done = 0
    for idx, row in df.iterrows():
        cfg_path = row.get("cfg_path")
        cfg_folder = row.get("cfg_folder")
        if not cfg_path or cfg_path == "-" or not cfg_folder or cfg_folder == "-":
            continue
        old_path = Path(cfg_path)
        new_path = old_path.parent.parent / cfg_folder / old_path.name
        if new_path == old_path:
            continue
        if not old_path.exists():
            log.debug("Skipping missing config file: %s", old_path)
            continue
        cfg = _read_cfg_file(old_path, defaults_file=find_default_files(old_path, stop_at=cfg_root))
        if "DATA" not in cfg.sections():
            continue
        done += 1
        if progress_callback is not None:
            progress_callback(done, total, f"Renaming {old_path.name}...")
        log.info("renaming %s to %s", old_path, new_path)
        new_path.parent.mkdir(parents=True, exist_ok=True)
        o = subprocess.run(
            ["git", "-C", str(old_path.parent), "mv", old_path.as_posix(),
             new_path.as_posix()],
            capture_output=True)
        if b"fatal" in o.stderr:
            os.rename(old_path, new_path)

        old_fold = old_path.parent.name
        new_fold = new_path.parent.name
        if old_fold != new_fold:
            rename_movie_filename(new_path, old_fold, new_fold)
        dfo.loc[idx, "cfg_path"] = new_path.as_posix()

    write_config_map(content_path, dfo)
    return dfo
