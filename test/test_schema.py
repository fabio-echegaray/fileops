import json
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from fileops.scripts._schema import (
    CFG_CONTENT_SCHEMA_VERSION,
    SCHEMA_CFG_CONTENT,
    SCHEMA_SUMMARY,
    SUMMARY_SCHEMA_VERSION,
    LegacyFormatError,
    detect_cfg_content_schema,
    detect_summary_schema,
    read_workbook_schema,
    set_workbook_schema,
)


class TestWorkbookSchema(unittest.TestCase):

    def setUp(self):
        self._tmpdir = Path(tempfile.mkdtemp())

    def _save(self, book, name="book.xlsx"):
        path = self._tmpdir / name
        book.save(path)
        return path

    def test_set_workbook_schema_stamps_document_properties(self):
        book = openpyxl.Workbook()
        book.active.title = "Files-Timeseries"
        set_workbook_schema(book, SCHEMA_SUMMARY, SUMMARY_SCHEMA_VERSION)

        payload = json.loads(book.properties.description)
        self.assertEqual(payload["schema"], SCHEMA_SUMMARY)
        self.assertEqual(payload["schema_version"], SUMMARY_SCHEMA_VERSION)

    def test_read_workbook_schema_roundtrip_via_disk(self):
        book = openpyxl.Workbook()
        book.active.title = "Files-Timeseries"
        set_workbook_schema(book, SCHEMA_SUMMARY, SUMMARY_SCHEMA_VERSION)
        path = self._save(book)

        self.assertEqual(read_workbook_schema(path),
                         (SCHEMA_SUMMARY, SUMMARY_SCHEMA_VERSION))

    def test_read_workbook_schema_returns_none_for_unmarked_book(self):
        book = openpyxl.Workbook()
        book.active.title = "Files-Timeseries"
        path = self._save(book)

        self.assertIsNone(read_workbook_schema(path))


class TestDetectSummarySchema(unittest.TestCase):

    def setUp(self):
        self._tmpdir = Path(tempfile.mkdtemp())

    def _make_summary(self, name, stamp_schema=False, with_cfg_columns=False):
        path = self._tmpdir / name
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame({"name": ["ch1"]}).to_excel(writer, sheet_name="Channels", index=False)
            columns = ["folder", "filename", "frames", "image_series_id"]
            if with_cfg_columns:
                columns += ["cfg_path", "cfg_folder"]
            pd.DataFrame(data=columns).T.to_excel(
                writer, sheet_name="Files-Timeseries", index=False, header=False)
        if stamp_schema:
            book = openpyxl.load_workbook(path)
            set_workbook_schema(book, SCHEMA_SUMMARY, SUMMARY_SCHEMA_VERSION)
            book.save(path)
        return path

    def test_detect_new_summary_from_schema_marker(self):
        path = self._make_summary("new.xlsx", stamp_schema=True)
        self.assertEqual(detect_summary_schema(path),
                         (SCHEMA_SUMMARY, SUMMARY_SCHEMA_VERSION))

    def test_detect_new_summary_with_stripped_properties_falls_back_to_structure(self):
        # an external save may strip the document properties; the structure
        # alone (no cfg columns) must still identify the current format
        path = self._make_summary("stripped.xlsx", stamp_schema=False)
        self.assertEqual(detect_summary_schema(path),
                         (SCHEMA_SUMMARY, SUMMARY_SCHEMA_VERSION))

    def test_detect_legacy_summary_raises(self):
        # legacy summaries carry cfg_path/cfg_folder columns and no marker;
        # they must hard-fail instead of being silently misread
        path = self._make_summary("legacy.xlsx", stamp_schema=False, with_cfg_columns=True)
        with self.assertRaises(LegacyFormatError):
            detect_summary_schema(path)

    def test_detect_non_xlsx_summary_raises(self):
        path = self._tmpdir / "summary.csv"
        path.write_text("folder,filename\n")
        with self.assertRaises(LegacyFormatError):
            detect_summary_schema(path)

    def test_detect_missing_summary_raises(self):
        path = self._tmpdir / "does-not-exist.xlsx"
        with self.assertRaises(LegacyFormatError):
            detect_summary_schema(path)


class TestDetectCfgContentSchema(unittest.TestCase):

    def setUp(self):
        self._tmpdir = Path(tempfile.mkdtemp())

    def _make_content(self, name, stamp_schema=False, with_configs_sheet=True):
        path = self._tmpdir / name
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame({
                "cfg_path": ["a.cfg"],
                "title":    ["t"],
            }).to_excel(writer, sheet_name="Content", index=False)
            if with_configs_sheet:
                pd.DataFrame({
                    "cfg_folder": [""],
                    "cfg_path":   [""],
                }).to_excel(writer, sheet_name="Configs", index=False)
        if stamp_schema:
            book = openpyxl.load_workbook(path)
            set_workbook_schema(book, SCHEMA_CFG_CONTENT, CFG_CONTENT_SCHEMA_VERSION)
            book.save(path)
        return path

    def test_detect_new_content_from_schema_marker(self):
        path = self._make_content("new.xlsx", stamp_schema=True)
        self.assertEqual(detect_cfg_content_schema(path),
                         (SCHEMA_CFG_CONTENT, CFG_CONTENT_SCHEMA_VERSION))

    def test_detect_new_content_with_stripped_properties_falls_back_to_structure(self):
        # no marker but the Configs sheet is present: still the current format
        path = self._make_content("stripped.xlsx", stamp_schema=False, with_configs_sheet=True)
        self.assertEqual(detect_cfg_content_schema(path),
                         (SCHEMA_CFG_CONTENT, CFG_CONTENT_SCHEMA_VERSION))

    def test_detect_legacy_content_raises(self):
        # legacy cfg_content files have only the data sheet, no Configs sheet
        # and no marker; they must hard-fail
        path = self._make_content("legacy.xlsx", stamp_schema=False, with_configs_sheet=False)
        with self.assertRaises(LegacyFormatError):
            detect_cfg_content_schema(path)

    def test_detect_non_xlsx_content_raises(self):
        path = self._tmpdir / "cfg_content.txt"
        path.write_text("cfg_path,title\n")
        with self.assertRaises(LegacyFormatError):
            detect_cfg_content_schema(path)


if __name__ == '__main__':
    unittest.main()